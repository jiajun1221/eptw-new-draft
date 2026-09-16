from pathlib import Path
from docx import Document
import pdfplumber

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "inputs"
OUT = ROOT / "extract"
OUT.mkdir(parents=True, exist_ok=True)

for name in ["database-redesign-proposal.docx", "old-vs-new-schema.docx"]:
    doc = Document(INPUT / name)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(f"[{p.style.name}] {p.text}")
    for ti, table in enumerate(doc.tables, 1):
        lines.append(f"\n[TABLE {ti}]")
        for row in table.rows:
            lines.append("\t".join(cell.text.replace("\n", " / ") for cell in row.cells))
    (OUT / f"{Path(name).stem}.txt").write_text("\n".join(lines), encoding="utf-8")

with pdfplumber.open(INPUT / "ePTW-Requirements-Review.pdf") as pdf:
    pages = []
    for i, page in enumerate(pdf.pages, 1):
        pages.append(f"\n===== PAGE {i} =====\n{page.extract_text() or ''}")
    (OUT / "ePTW-Requirements-Review.txt").write_text("\n".join(pages), encoding="utf-8")

